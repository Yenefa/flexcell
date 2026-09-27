# -*- coding: utf-8 -*-
"""Generate Day-2 acceptance report docx for project-033-flexcell (for 大统领)."""
import os

from docx import Document
from docx.shared import Pt

OUT = r"C:\Users\fuker\Desktop\workspace\project-033-flexcell\docs\验收报告-project-033-flexcell-day2-20260926.docx"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)

doc.add_heading("project-033-flexcell Day-2 验收报告 · MoveToPose 技能接口", level=0)
p = doc.add_paragraph("日期：2026-09-26　|　上游：Yenefa/flexcell @ a6b00e7　|　执行环境：本机 WSL2 Ubuntu-24.04（user ros）")
p = doc.add_paragraph("结论速览：上游 Day-2《MoveToPose ROS 2 Action》在本机全量复验，10/10 用例通过；技能服务器随仿真栈常驻，复验命令见第五节。")

doc.add_heading("一、验收结论", level=1)
doc.add_paragraph(
    "Day-1 打好的环境底座上，Day-2 的目标是把「人看着 RViz 点 Plan/Execute」升级为「agent 发一个类型化请求 → 服务器返回带实测验证的类型化结果」。"
    "本次在本机完整复现上游声明的 10 个验收用例，全部通过：两条 happy path 的成功是 TF 实测出来的（位姿误差 6.1mm / 8.2mm，阈值 20mm），"
    "五类非法输入全部被正确拒绝且绝不触碰执行器，超时与取消都能真实打断运动并给出严格确认（preempted + stop_confirmed），Day-1 能力回归无损。"
)

doc.add_heading("二、Day-2 DoD 验收表（10/10）", level=1)
rows = [
    ("#", "用例", "输入要点", "实测结果", "判定"),
    ("1", "happy path A", "目标 (0.20,0.05,0.20)，demo_slow", "SUCCEEDED，实测位姿误差 0.0061 m，四阶段反馈齐全", "PASS"),
    ("2", "happy path B", "目标 (0.15,-0.15,0.25)", "SUCCEEDED，实测位姿误差 0.0082 m", "PASS"),
    ("3", "工作区外", "目标 z=0.90（上限 0.75）", "INVALID_GOAL，38ms 拒绝，未碰执行器", "PASS"),
    ("4", "四元数未归一", "q=(0,0,0,0.5)", "INVALID_GOAL，detail 报 norm=0.5", "PASS"),
    ("5", "画像不在白名单", "profile=turbo", "INVALID_GOAL，detail 列出合法画像", "PASS"),
    ("6", "世界版本闸门", "expected_world_revision=7（闸门关闭）", "STALE_WORLD，拒绝造假而非假装通过", "PASS"),
    ("7", "IK 不可达", "目标 (0.42,0,0.10)（箱内臂展外）", "IK_UNREACHABLE，1.06s 快败，execution_started=false", "PASS"),
    ("8", "执行超时", "deadline=800ms + 大摆动靶点", "EXECUTION_TIMEOUT，812ms 进 STOPPING，preempted=true，stop_confirmed=true", "PASS"),
    ("9", "执行中取消", "EXECUTING 1.0s 后 cancel", "CANCELED，preempted=true，stop_confirmed=true", "PASS"),
    ("10", "Day-1 回归", "plan_execute.py 0.0 -0.9 0.9 0.5", "PLAN/EXECUTE error_code=1，关节1 实摆 1.03 rad，MOTION_DETECTED=True", "PASS"),
]
t = doc.add_table(rows=len(rows), cols=5)
t.style = "Table Grid"
for i, r in enumerate(rows):
    for j, cell in enumerate(r):
        t.rows[i].cells[j].text = cell
doc.add_paragraph("注：case8 首跑靶点 (0.30,0.25,0.15) 球坐标约 0.41m 超出臂展（~0.38m），被 IK 预检快败——属测试靶点设计问题而非服务器缺陷，换靶点后通过。")

doc.add_heading("三、环境与资产", level=1)
for line in [
    "构建：/home/ros/agentic_workcell_ws 内 colcon build --packages-select flexcell_interfaces flexcell_skills → 2 包 finished [29.6s]。",
    "运行：/move_to_pose 动作服务器（flexcell_skills/move_to_pose.launch.py，controller_action 后端，pos_tol=0.020m，orientation_enforced=false 显式上报）。",
    "证据：F:\\flexcell-stage\\setup\\acceptance\\day2\\（00-11 共 12 个输出 + move_to_pose.jsonl 结构化日志 10 条 + 服务器日志）。",
    "验收电池脚本：F:\\flexcell-stage\\setup\\32_day2_acceptance.sh（可重复执行，用例名传参）；归档脚本 33_day2_archive.sh。",
    "上游仓库：本地克隆与 WSL 工作区均在 a6b00e7，工作树 clean，无本地改动，未 push。",
]:
    doc.add_paragraph(line, style="List Bullet")

doc.add_heading("四、技能接口与关键行为（给 agent 的契约）", level=1)
for line in [
    "调用方式：/move_to_pose 动作（flexcell_interfaces/action/MoveToPose），请求含 request_id / target(PoseStamped) / motion_profile_id / expected_world_revision / deadline_ms。",
    "结果含：outcome（SUCCEEDED/FAILED/CANCELED/STATE_UNKNOWN）+ 错误码分类 + MoveIt 原始码 + TF 实测 reached_pose 与位置/朝向误差 + execution_started + 抢占/停止确认 + evidence_id。",
    "成功是测出来的：MoveIt 报成功不算成功，必须 TF 实测进容差；朝向不做硬约束（上游 KDL 配置 position_only_ik），但每次结果都显式声明 orientation_enforced=false，绝不静默放宽。",
    "可打断性：执行走 controller_action 后端（FlexCell 直握 /arm_controller/follow_joint_trajectory 目标），取消/超时实测 ~0.05-1s 内停住；MoveIt 2.12.4 自带 /execute_trajectory 的取消是坏的（上游已调查定案），仅留作审计。",
    "每次请求落一条 JSONL 结构化日志（~/.ros/flexcell/move_to_pose.jsonl），含验证/规划/执行/停止全过程与 evidence_id，可审计可回放。",
]:
    doc.add_paragraph(line, style="List Bullet")

doc.add_heading("五、复验步骤（大统领操作）", level=1)
for line in [
    "1) 查栈与技能服务器是否存活（预期 gz sim/move_group/rviz2 + move_to_pose_server）：",
    "   wsl -d Ubuntu-24.04 -u ros -- pgrep -af \"move_group|gz sim|rviz2|move_to_pose_server\"",
    "2) 若栈已死，重启并验活：",
    "   wsl -d Ubuntu-24.04 -u ros -- bash /mnt/f/flexcell-stage/setup/50_start_detached.sh",
    "   wsl -d Ubuntu-24.04 -u ros -- pgrep -af move_group",
    "3) 若技能服务器未活，单独拉起（命令见 project-033-flexcell README 快速开始 2.5 步）。",
    "4) 跑 Day-2 全量验收电池（约 2 分钟，10 用例）：",
    "   wsl -d Ubuntu-24.04 -u ros -- bash -c 'tr -d \"\\r\" < /mnt/f/flexcell-stage/setup/32_day2_acceptance.sh > /tmp/d2.sh && bash /tmp/d2.sh happy_a happy_b invalid_workspace invalid_quat invalid_profile stale_world ik_unreachable timeout cancel regression'",
    "5) 判定：11_summary.out 汇总 10 行全 PASS（happy 两条看 position_error_m < 0.02；regression 看 MOTION_DETECTED=True——若上一步刚跑过回归臂已在靶位，DELTA≈0 属正常）。",
]:
    doc.add_paragraph(line)

doc.add_heading("六、本次踩坑摘要（详见 踩坑日记/开发踩坑.md）", level=1)
for line in [
    "wsl.exe interop 撕复杂内联命令：多行/嵌套引号/转义会被截断重排，裸 /tmp 路径还会被换成 Windows Temp——复杂命令一律写脚本文件（tr -d 去 CR）再执行，脚本本身沉淀为可复现资产。",
    "set -u 与 ROS setup.bash 冲突（AMENT_TRACE_SETUP_FILES unbound variable），source ROS 的脚本禁用 set -u。",
    "验收负路径用例的靶点要先按臂展包络（~0.35m 内）验算再写死，否则被 IK_UNREACHABLE 顶替目标用例。",
]:
    doc.add_paragraph(line, style="List Bullet")

doc.save(OUT)
print("saved:", OUT, os.path.getsize(OUT), "bytes")
